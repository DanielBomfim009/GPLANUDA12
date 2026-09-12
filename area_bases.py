"""Área de Bases: carregar uma base por vez, relacionar as colunas do arquivo
com o que o Gplan usa e refazer só o que depende dela.

Fica fora do gplan_app.py porque não desenha nada: é o que mexe na pasta das
bases, grava o mapeamento que o pipeline lê e roda o pipeline. A tela
(render_bases, no app) só chama o que está aqui.

O mapeamento de cada base vai para
00_BASES_ATUALIZACAO/02_MODELOS_E_INSTRUCOES/mapeamentos/<arquivo>.json, e quem
o aplica é o load_sheet_rows do pipeline (update_current_workbook_from_bases.py):
ele renomeia as colunas do arquivo para os nomes que as leituras já procuram.
O nome "padrão" de cada campo aqui é o PRIMEIRO apelido que a leitura
correspondente do pipeline usa -- se lá mudar, muda aqui.

A atualização roda numa thread, fora da execução da página. Leva de um a
alguns minutos, e o Streamlit refaz a página a cada clique: se o pipeline
rodasse dentro dela, trocar de aba no meio interromperia a execução antes de
desfazer uma troca que falhou. A página só acompanha (tarefa_atual).

Caminhos: GPLAN_CONTROLE_DIR (a pasta "Controle de Relatório dos
Instrumentos", de onde vêm os scripts) e GPLAN_BASES_DIR, GPLAN_BACKUPS e
GPLAN_LOGS (as mesmas variáveis que o pipeline lê). Sem elas, valem as pastas
ao lado do app, como no LOCAL_EXCEL_FALLBACK. A instância de desenvolvimento
aponta bases, backups e logs para uma cópia: os scripts são os de verdade, e
nenhum teste toca nas bases nem na planilha de verdade.
"""
from __future__ import annotations

import datetime as dt
import difflib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import unicodedata
import uuid
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

APP_DIR = Path(__file__).resolve().parent
VERSOES_GUARDADAS = 5
RESUMO_ABAS = "07_TAG_RESUMO · 08_RELATORIOS_ESPERADOS · 09_PENDENCIAS · 12_CALC_*"
PLANILHA_SUPABASE = "CONTROLE_DOCUMENTAL_INSTRUMENTACAO_ATUAL.xlsx"


# ------------------------------------------------------------------ caminhos
def controle_dir() -> Path:
    v = os.environ.get("GPLAN_CONTROLE_DIR")
    return Path(v) if v else APP_DIR.parent / "Controle de Relatório dos Instrumentos"


def bases_dir() -> Path:
    v = os.environ.get("GPLAN_BASES_DIR")
    return Path(v) if v else controle_dir() / "00_BASES_ATUALIZACAO"


def entrada_dir() -> Path:
    return bases_dir() / "00_COLOCAR_ATUALIZADAS_AQUI"


def historico_dir() -> Path:
    return bases_dir() / "03_HISTORICO_IMPORTACOES"


def mapeamentos_dir() -> Path:
    return bases_dir() / "02_MODELOS_E_INSTRUCOES" / "mapeamentos"


def ferramentas_dir() -> Path:
    return controle_dir() / "90_AREA_TECNICA" / "workbook_tools"


def recalcular_ps1() -> Path:
    return controle_dir() / "90_AREA_TECNICA" / "RECALCULAR_BASES.ps1"


def logs_dir() -> Path:
    v = os.environ.get("GPLAN_LOGS")
    return Path(v) if v else controle_dir() / "04_HISTORICO_VERSOES_E_TESTES" / "logs"


def disponivel() -> bool:
    """Só existe onde a pasta das bases e o pipeline existem: no computador de
    quem mantém as bases. No servidor (Render) não há pasta nenhuma."""
    return (entrada_dir().is_dir()
            and (ferramentas_dir() / "update_current_workbook_from_bases.py").is_file())


def normaliza(valor: object) -> str:
    """A mesma normalização do normalize_key do pipeline: sem acento, sem
    quebra de linha, maiúsculas, espaços colapsados."""
    texto = str(valor or "").strip().upper().replace("\n", " ")
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return " ".join(texto.split())


def _nome(texto: str) -> str:
    """Nome de arquivo comparável: sem diferença de maiúscula (a 04 é .XLSX)
    e com o Ç na mesma forma Unicode -- o OneDrive às vezes grava decomposto."""
    return unicodedata.normalize("NFC", str(texto)).lower()


# ------------------------------------------------------------------ cadastro
@dataclass(frozen=True)
class Campo:
    padrao: str                      # o nome que a leitura do pipeline procura
    rotulo: str
    obrigatorio: bool = False
    dica: str = ""
    apelidos: tuple = ()             # outros nomes que a leitura também aceita


@dataclass(frozen=True)
class Base:
    codigo: str
    nome: str
    arquivo: str
    extensoes: tuple                 # na ordem em que o pipeline procura
    unidade: str
    abas_sistema: tuple              # abas do controle que são dela (a 1ª dá a contagem)
    refaz: str
    impacto: str                     # "principal" | "resumo" | "leve"
    precisa_resumo: bool             # dispara o passo 3 (resumo por TAG)
    aba: str | None = None           # aba que o pipeline lê (None = a primeira)
    aba_fixa: bool = False           # o pipeline lê sempre esta aba, pelo nome
    cabecalho: int = 1               # linha do cabeçalho que o pipeline usa
    cabecalho_fixo: bool = False     # o pipeline procura o cabeçalho sempre nesta linha
    campos: tuple = ()
    chave: tuple = ()                # campos que identificam a linha (para comparar versões)
    fixa: str = ""                   # por que não tem mapeamento (leitura por posição)
    abas_exigidas: tuple = ()
    abas_refeitas: tuple = ()        # o que o passo 2 regrava; vazio = as abas_sistema

    @property
    def stem(self) -> str:
        return Path(self.arquivo).stem

    @property
    def mapeavel(self) -> bool:
        return bool(self.campos)

    @property
    def refeitas(self) -> tuple:
        return self.abas_refeitas or self.abas_sistema


def _c(padrao, rotulo, obrigatorio=False, dica="", *apelidos):
    return Campo(padrao, rotulo, obrigatorio, dica, tuple(apelidos))


BASES: tuple[Base, ...] = (
    Base("01", "TAGs", "01_BASE_TAGS.xlsx", (".xlsx",), "TAGs", ("01_BASE_TAGS",),
         "TAGs, Cabos, Locação, GITEC, de-para e o resumo por TAG", "principal", True,
         chave=("TAG",),
         abas_refeitas=("01_BASE_TAGS", "02_BASE_CABOS", "05_BASE_LOCAÇÃO", "05_AUX_AREAS",
                        "06_BASE_GITEC", "02_CABOS_DEPARA"),
         campos=(
             _c("TAG", "TAG", True, "identifica o instrumento"),
             _c("DESCRICAO", "Descrição", False, "sem ela, vale a coluna B", "DESCRIÇÃO"),
             _c("COMUNICAÇÃO", "Comunicação"),
             _c("TIPO_ORIGEM", "Tipo de origem", False, "sem ela, vale a posição de hoje"),
             _c("ITEM_PPU", "Item de PPU", False, "sem ela, vale a posição de hoje"),
             _c("FASE", "Fase"), _c("SOP", "SOP"), _c("SSOP", "SSOP"),
             _c("SUBGRUPO DE PRIORIDADE", "Subgrupo de prioridade"),
             _c("SEGMENTO", "Segmento"), _c("MALHA", "Malha"),
             _c("CRITÉRIO DE MEDIÇÃO ANEXO III - APÊNDICE I", "Critério de medição", False, "",
                "CRITERIO DE MEDICAO"),
             _c("PREÇO UNITÁRIO", "Preço unitário", False, "", "PRECO UNITARIO"),
             _c("STATUS DE LOCALIZAÇÃO", "Status de localização"),
             _c("STATUS DE CALIBRAÇÃO", "Status de calibração"),
             _c("STATUS DE MONTAGEM", "Status de montagem"),
             _c("STATUS FINAL", "Status final"),
             _c("CFF", "Caixa (CFF)"), _c("PAINEL", "Painel"),
             _c("SSOP PRIORITÁRIO", "SSOP prioritário"), _c("SKID", "Skid"),
             _c("SEMANA PROGRAMADA", "Semana programada"),
             _c("PREVISÃO TESTE MALHA", "Previsão do teste de malha"),
             _c("FORNECIMENTO", "Fornecimento (obrigação)"),
         )),
    Base("02", "Cabos", "02_BASE_CABOS.xlsx", (".xlsx",), "linhas", ("02_BASE_CABOS",),
         "a própria aba e o resumo por TAG", "resumo", True, cabecalho_fixo=True,
         chave=("TAG", "CABO"), campos=(
             _c("TAG", "TAG", True), _c("CABO", "Cabo", True),
             _c("ORIGEM", "Origem", True), _c("DESTINO", "Destino", True),
         )),
    Base("03", "Tubing", "03_BASE_TUBING.xlsx", (".xlsx",), "linhas", ("03_BASE_TUBING",),
         "a própria aba e o resumo por TAG", "resumo", True, chave=("TAG",), campos=(
             _c("TAG", "TAG", True),
             _c("TAG_CIRCUITO", "TAG do circuito", True, "", "TAG CIRCUITO"),
             _c("ORIGEM", "Origem", True),
         )),
    Base("04", "SIGEM", "04_BASE_SIGEM.xlsx", (".xlsx",), "documentos", ("04_BASE_SIGEM",),
         "a própria aba e o resumo por TAG", "resumo", True, chave=("DOCUMENTO", "REVISAO"),
         campos=(
             _c("DOCUMENTO", "Documento", True),
             _c("STATUS", "Status", True),
             _c("REVISAO", "Revisão", True, "", "REVISÃO"),
             _c("DATA", "Data de inclusão", True, "", "Incluido em"),
             _c("TITULO", "Título", True, "", "TÍTULO"),
         )),
    Base("05", "Locação", "05_BASE_LOCAÇÃO.xlsx", (".xlsx",), "locações",
         ("05_BASE_LOCAÇÃO", "05_AUX_AREAS"),
         "a própria aba, as áreas e o resumo por TAG", "resumo", True, chave=("TAG",), campos=(
             _c("TAG", "TAG", True),
             _c("DESCRICAO", "Descrição", True, "", "DESCRIÇÃO"),
             _c("LOCACAO", "Locação", True, "", "LOCAÇÃO"),
             _c("BANDEJA", "Bandeja", True), _c("ELETRODUTO", "Eletroduto", True),
             _c("SUPORTE", "Suporte", True),
             _c("AREA", "Área", False, "sem ela, a aba Planta fica sem número", "ÁREA"),
         )),
    Base("06", "GITEC", "06_BASE_GITEC.xlsx", (".xlsx",), "medições", ("06_BASE_GITEC",),
         "a própria aba e o resumo por TAG", "resumo", True, aba="RelResumoEvento",
         aba_fixa=True, chave=("TAG",), campos=(
             _c("TAG", "TAG", True), _c("FASE", "Fase", True),
             _c("AGRUPAMENTO", "Agrupamento", True, "de onde sai o item de PPU"),
             _c("ETAPA", "Etapa", True), _c("STATUS", "Status", True),
             _c("VALOR", "Valor", True),
             _c("DATA DE EXECUCAO", "Data de execução", True, "", "DATA DE EXECUÇÃO"),
         )),
    Base("07", "Cabos completo", "07_BASE_CABOS_COMPLETO.xlsm", (".xlsm", ".xlsx"), "circuitos",
         ("02_CABOS_LANCAMENTO", "02_CABOS_DEPARA"),
         "lançamento, de-para e o resumo por TAG", "resumo", True,
         aba="BASE DE DADOS", aba_fixa=True, cabecalho=13, chave=("CIRCUITO",), campos=(
             _c("CIRCUITO", "Circuito", True, "chave de cada cabo"),
             _c("DE", "Origem", True, "ponta de campo"),
             _c("PARA", "Destino", True, "caixa ou painel"),
             _c("DISCIPLINA", "Disciplina", True),
             _c("TIPO", "Tipo", True),
             _c("STATUS LANÇAMENTO", "Status do lançamento", True),
             _c("% de Conclusão LANÇAMENTO", "% de lançamento", True),
             _c("COMP. TOTAL DO CABO (m)", "Metragem prevista", True),
             _c("COMPR.(M) CAMPO REALIZADO", "Metragem lançada", False, "o que o campo mediu"),
             _c("% de Conclusão CONEXÃO", "% de conexão"),
             _c("1ª Ponta (Status)", "1ª ponta", False, "data da conexão"),
             _c("2ª Ponta (Status)", "2ª ponta", False, "data da conexão"),
             _c("Teste (Status)", "Teste", False, "data do teste"),
             _c("DOCUMENTO REF.", "Documento de referência"),
             _c("% Avanço REAL", "% avanço real", False, "lançamento + conexão + teste"),
         )),
    Base("08", "Pedestal", "08_BASE_PEDESTAL.xlsx", (".xlsx",), "pedestais", ("08_BASE_PEDESTAL",),
         "a própria aba e o resumo por TAG", "resumo", True, cabecalho=2,
         chave=("RIR DO SUPORTE (MEDIÇÃO)",), campos=(
             _c("RIR DO SUPORTE (MEDIÇÃO)", "RIR do suporte", True, "documento do pedestal",
                "RIR DO SUPORTE (MEDICAO)"),
             *[_c(f"TAGINSTR{n}", f"TAG {n}") for n in range(1, 13)],
             _c("AVANÇO", "Avanço", False, "", "AVANCO"),
         )),
    Base("09", "Suprimentos", "09_BASE_SUPRIMENTOS.xlsx", (".xlsx",), "itens",
         ("09_SUPRIMENTOS_ITENS", "09_SUPRIMENTOS_ESTOQUE"),
         "itens, estoque, as colunas de fornecimento das TAGs e o resumo", "resumo", True,
         aba="Mapa de Suprimentos UDA", cabecalho=2,
         abas_refeitas=("01_BASE_TAGS", "09_SUPRIMENTOS_ITENS", "09_SUPRIMENTOS_ESTOQUE"),
         fixa=("Lida por posição no Mapa de Suprimentos UDA: duas linhas de cabeçalho e "
               "cerca de 100 colunas fixas. Aqui a tela confere o leiaute em vez de "
               "relacionar colunas."),
         abas_exigidas=("Mapa de Suprimentos UDA", "almoxarifado-estoque-smat")),
    Base("10", "Rundown", "10_BASE_RUNDOWN.xlsx", (".xlsx",), "semanas",
         ("10_BASE_RUNDOWN_CURVA", "10_BASE_RUNDOWN", "10_BASE_RUNDOWN_CAL"),
         "só a própria aba", "leve", False, aba="CURVA", cabecalho=2,
         fixa="Arquivo de configuração do Rundown, com PARAMETROS, CURVA e CALENDARIO fixos.",
         abas_exigidas=("PARAMETROS", "CURVA", "CALENDARIO")),
    Base("11", "Infraestrutura", "11_BASE_INFRAESTRUTURA.xlsx", (".xlsx",), "desenhos",
         ("11_BASE_INFRAESTRUTURA",), "só a própria aba", "leve", False,
         chave=("DESENHO",), campos=(
             _c("DESENHO", "Desenho", True),
             _c("AVANÇO GERAL", "Avanço geral", True, "", "AVANCO GERAL"),
         )),
)
POR_CODIGO = {b.codigo: b for b in BASES}


# --------------------------------------------------------- estado da pasta
def arquivo_atual(b: Base) -> Path | None:
    """O arquivo que o pipeline vai ler: o primeiro que existir na ordem das
    extensões (a 07 aceita .xlsm e .xlsx, e o pipeline prefere o .xlsm)."""
    pasta = entrada_dir()
    if not pasta.is_dir():
        return None
    por_nome = {_nome(p.name): p for p in pasta.iterdir() if p.is_file()}
    for ext in b.extensoes:
        achado = por_nome.get(_nome(b.stem + ext))
        if achado:
            return achado
    return None


def mapeamento_salvo(b: Base) -> dict:
    arq = mapeamentos_dir() / f"{b.stem}.json"
    if not arq.exists():
        return {}
    try:
        return json.loads(arq.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def versoes(b: Base) -> list[Path]:
    pasta = historico_dir() / b.stem
    if not pasta.is_dir():
        return []
    return sorted((p for p in pasta.iterdir() if p.is_file()), key=lambda p: p.name, reverse=True)


def _ler_fim(arq: Path, limite: int = 400_000) -> list[str]:
    """As últimas linhas de um log sem ler o arquivo inteiro."""
    with open(arq, "rb") as f:
        f.seek(0, 2)
        tamanho = f.tell()
        f.seek(max(0, tamanho - limite))
        texto = f.read().decode("utf-8", errors="replace")
    linhas = texto.splitlines()
    return linhas[1:] if tamanho > limite else linhas    # a primeira pode vir cortada


_CARIMBO = re.compile(r"\[(\d{4}-\d\d-\d\d \d\d:\d\d:\d\d)\]")


def ultima_execucao() -> dict:
    """A última atualização completa e a última individual, do log do
    pipeline (importar_bases.log). O fim do resumo por TAG (aplicar_formulas.log)
    conta para a importação que terminou logo antes dele, e só se nenhuma
    outra começou no meio -- senão um resumo rodado depois de uma carga
    individual esticava a duração da última completa."""
    rodadas = []
    arq = logs_dir() / "importar_bases.log"
    if arq.exists():
        atual = None
        for linha in _ler_fim(arq):
            m = _CARIMBO.match(linha)
            if not m:
                continue
            quando = dt.datetime.strptime(m.group(1), "%Y-%m-%d %H:%M:%S")
            if "Iniciando importa" in linha:
                atual = {"inicio": quando, "bases": None, "fim": None, "fim_resumo": None}
            elif atual and "Atualizacao individual" in linha:
                achado = re.search(r"base\(s\) ([\d, ]+)", linha)
                atual["bases"] = achado.group(1).strip() if achado else "?"
            elif atual and "Importação concluída" in linha:
                atual["fim"] = quando
                rodadas.append(atual)
                atual = None
    fins = []
    resumo = logs_dir() / "aplicar_formulas.log"
    if resumo.exists():
        for linha in _ler_fim(resumo, 150_000):
            m = _CARIMBO.match(linha)
            if m and "Atualização concluída" in linha:
                fins.append(dt.datetime.strptime(m.group(1), "%Y-%m-%d %H:%M:%S"))
    for i, r in enumerate(rodadas):
        seguinte = rodadas[i + 1]["inicio"] if i + 1 < len(rodadas) else None
        depois = [f for f in fins if r["fim"] <= f <= r["fim"] + dt.timedelta(minutes=20)
                  and (seguinte is None or f < seguinte)]
        if depois:
            r["fim_resumo"] = depois[0]
    return {"completa": next((r for r in reversed(rodadas) if not r["bases"]), None),
            "individual": next((r for r in reversed(rodadas) if r["bases"]), None)}


_CONTAGENS: dict = {}


def contagens(planilha: Path) -> dict[str, int]:
    """Linhas de cada aba do controle, para o cartão de cada base. Guardado
    pelo carimbo da planilha: só relê quando ela muda."""
    try:
        carimbo = (str(planilha), planilha.stat().st_mtime)
    except OSError:
        return {}
    if _CONTAGENS.get("carimbo") != carimbo:
        import openpyxl
        valores = {}
        try:
            wb = openpyxl.load_workbook(planilha, read_only=True)
        except Exception:
            return {}
        try:
            for aba in {b.abas_sistema[0] for b in BASES}:
                if aba in wb.sheetnames:
                    ws = wb[aba]
                    total = ws.max_row or sum(1 for _ in ws.iter_rows(values_only=True))
                    valores[aba] = max(0, total - 1)
        finally:
            wb.close()
        _CONTAGENS.update(carimbo=carimbo, valores=valores)
    return _CONTAGENS["valores"]


# ----------------------------------------------------------- leitura do arquivo
@dataclass
class Analise:
    abas: list
    aba: str = ""
    cabecalho: int = 1
    cabecalho_como: str = ""        # detectado | salvo | escolhido | fixo | padrão
    colunas: list = field(default_factory=list)
    amostras: dict = field(default_factory=dict)   # coluna -> 3 primeiros valores
    primeiras: list = field(default_factory=list)  # [(nº da linha, valores)] das 12 primeiras
    linhas: int = 0
    formulas: int = 0               # fórmulas nas primeiras linhas de dado
    sem_valor: int = 0              # dessas, as que não têm valor salvo
    faltam_abas: list = field(default_factory=list)
    erro: str = ""


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float):
        return f"{v:.4g}"
    return " ".join(str(v).split())


def _nomes_conhecidos(b: Base, salvo: dict) -> set:
    nomes = {normaliza(n) for c in b.campos for n in (c.padrao, *c.apelidos)}
    for dados in (salvo.get("abas") or {}).values():
        nomes |= {normaliza(k) for k in (dados.get("colunas") or {})}
    return nomes


def _pontua(nomes: set, linha) -> int:
    return sum(1 for v in (linha or ()) if v is not None and normaliza(v) in nomes)


def _conta_linhas(ws, cabecalho: int) -> int:
    return sum(1 for r in ws.iter_rows(min_row=cabecalho + 1, values_only=True)
               if any(v not in (None, "") for v in r))


def analisar(b: Base, conteudo: bytes, aba: str | None = None,
             cabecalho: int | None = None) -> Analise:
    """O que tem dentro do arquivo: abas, a aba e a linha do cabeçalho mais
    prováveis, as colunas com três exemplos cada e quantas linhas de dado."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as erro:
        return Analise(abas=[], erro=("Não consegui abrir o arquivo como planilha do Excel "
                                      f"({type(erro).__name__}: {erro})."))
    try:
        abas = list(wb.sheetnames)
        if b.fixa:
            an = Analise(abas=abas, faltam_abas=[a for a in b.abas_exigidas if a not in abas])
            an.aba = b.aba if b.aba in abas else abas[0]
            an.cabecalho, an.cabecalho_como = b.cabecalho, "fixo"
            if not an.faltam_abas:
                an.linhas = _conta_linhas(wb[an.aba], b.cabecalho)
        else:
            an = _analisar_colunas(b, wb, abas, aba, cabecalho)
    finally:
        wb.close()
    if not an.erro and an.aba:
        an.formulas, an.sem_valor = _formulas(conteudo, an.aba, an.cabecalho)
    return an


def _analisar_colunas(b, wb, abas, aba, cabecalho) -> Analise:
    salvo = mapeamento_salvo(b)
    if b.aba_fixa:
        if b.aba not in abas:
            return Analise(abas=abas, erro=(
                f'O arquivo não tem a aba "{b.aba}", que é a que o Gplan lê nesta base. '
                f"Abas do arquivo: {', '.join(abas)}."))
        escolhida = b.aba
    else:
        escolhida = aba if aba in abas else None
        if not escolhida and salvo.get("aba_principal") in abas:
            escolhida = salvo["aba_principal"]
        if not escolhida and b.aba in abas:
            escolhida = b.aba
    nomes = _nomes_conhecidos(b, salvo)
    topos = {}
    if not escolhida:
        melhor = -1
        for nome in abas:
            topos[nome] = list(wb[nome].iter_rows(max_row=40, values_only=True))
            nota = max((_pontua(nomes, r) for r in topos[nome]), default=0)
            if nota > melhor:
                melhor, escolhida = nota, nome
    ws = wb[escolhida]
    topo = topos.get(escolhida) or list(ws.iter_rows(max_row=40, values_only=True))
    notas = [(_pontua(nomes, r), -i) for i, r in enumerate(topo)]
    nota_melhor, menos_i = max(notas) if notas else (0, 0)
    salvo_cab = ((salvo.get("abas") or {}).get(escolhida) or {}).get("cabecalho")
    if b.cabecalho_fixo:
        cab, como = b.cabecalho, "fixo"
    elif cabecalho:
        cab, como = int(cabecalho), "escolhido"
    elif salvo_cab and len(topo) >= int(salvo_cab) and \
            _pontua(nomes, topo[int(salvo_cab) - 1]) >= nota_melhor:
        cab, como = int(salvo_cab), "salvo"
    elif nota_melhor > 0:
        cab, como = -menos_i + 1, "detectado"
    else:
        cab, como = b.cabecalho, "padrão"
    if len(topo) >= cab:
        linha_cab = topo[cab - 1]
    else:
        linha_cab = next(ws.iter_rows(min_row=cab, max_row=cab, values_only=True), ())
    colunas = [_texto(v) for v in linha_cab]
    while colunas and not colunas[-1]:
        colunas.pop()
    amostras: dict[str, list] = {c: [] for c in colunas if c}
    primeiras = []
    linhas = 0
    for n, row in enumerate(ws.iter_rows(min_row=cab + 1, values_only=True), start=cab + 1):
        if not any(v not in (None, "") for v in row):
            continue
        linhas += 1
        if linhas <= 12:
            primeiras.append((n, [_texto(v) for v in row[:len(colunas)]]))
        if linhas <= 400:
            for j, v in enumerate(row[:len(colunas)]):
                c = colunas[j]
                if c and len(amostras[c]) < 3 and v not in (None, ""):
                    amostras[c].append(_texto(v))
    return Analise(abas=abas, aba=escolhida, cabecalho=cab, cabecalho_como=como,
                   colunas=colunas, amostras=amostras, primeiras=primeiras, linhas=linhas)


_NS_PLANILHA = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_NS_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _parte_da_aba(z: zipfile.ZipFile, aba: str) -> str | None:
    """O XML de uma aba dentro do .xlsx, pelo nome dela."""
    rid = next((s.get(_NS_REL + "id") for s in ET.fromstring(z.read("xl/workbook.xml"))
                .iter(_NS_PLANILHA + "sheet") if s.get("name") == aba), None)
    for rel in ET.fromstring(z.read("xl/_rels/workbook.xml.rels")):
        if rel.get("Id") == rid:
            alvo = rel.get("Target", "")
            return alvo.lstrip("/") if alvo.startswith("/") else "xl/" + alvo
    return None


def _formulas(conteudo: bytes, aba: str, cabecalho: int, limite: int = 300) -> tuple[int, int]:
    """Fórmulas nas primeiras linhas de dado, e quantas estão sem valor
    salvo. O pipeline lê o último valor que o Excel gravou: sem valor, a
    célula entra vazia. É o que obriga, hoje, a abrir as bases no Excel.

    Lido do XML, e não pelo openpyxl: para ele, "o Excel calculou e deu
    texto vazio" (t="str" com <v/>, caso do SKID da 01, que vem de vínculo
    externo) e "nunca foi calculada" (sem <v>) são o mesmo None -- e a 01
    aparecia com 268 fórmulas sem valor que na verdade estavam em dia."""
    try:
        z = zipfile.ZipFile(io.BytesIO(conteudo))
        parte = _parte_da_aba(z, aba)
        if not parte:
            return 0, 0
        formulas = sem_valor = 0
        with z.open(parte) as f:
            for _, el in ET.iterparse(f):
                if el.tag != _NS_PLANILHA + "row":
                    continue
                n = int(el.get("r") or 0)
                if n > cabecalho + limite:
                    break
                if n > cabecalho:
                    for c in el.iter(_NS_PLANILHA + "c"):
                        if c.find(_NS_PLANILHA + "f") is None:
                            continue
                        formulas += 1
                        v = c.find(_NS_PLANILHA + "v")
                        if v is None or (not v.text and c.get("t") != "str"):
                            sem_valor += 1
                el.clear()
        return formulas, sem_valor
    except Exception:
        return 0, 0


# ------------------------------------------------------------- colunas
def sugerir(b: Base, colunas: list[str]) -> list[dict]:
    """Para cada campo: a coluna sugerida e de onde veio a sugestão.

    salvo      o mapeamento da última carga desta pasta
    nome       o nome da coluna é um dos que o pipeline já aceita
    sugestao   só um nome parecido -- vale conferir
    nao_usar   na última carga, a coluna existia e foi marcada "não usar"
    (vazio)    nada encontrado
    """
    norm: dict[str, str] = {}
    for c in colunas:
        if c:
            norm.setdefault(normaliza(c), c)
    salvo = mapeamento_salvo(b)
    salvo_col: dict[str, str] = {}
    for dados in (salvo.get("abas") or {}).values():
        for arq, padrao in (dados.get("colunas") or {}).items():
            if not str(padrao).endswith("(ignorada)"):
                salvo_col.setdefault(normaliza(padrao), arq)
    nao_usar = {normaliza(p) for p in salvo.get("nao_usar") or []}
    tem_salvo = bool(salvo_col or nao_usar)
    saida = []
    for campo in b.campos:
        k = normaliza(campo.padrao)
        escolha, como, sumiu = None, "", ""
        if k in nao_usar:
            como = "nao_usar"
        else:
            arq = salvo_col.get(k)
            if arq and normaliza(arq) in norm:
                escolha, como = norm[normaliza(arq)], "salvo"
            elif arq:
                sumiu = arq
            if not escolha:
                for nome in (campo.padrao, *campo.apelidos):
                    if normaliza(nome) in norm:
                        escolha, como = norm[normaliza(nome)], "nome"
                        break
            if not escolha:
                perto = difflib.get_close_matches(k, list(norm), n=1, cutoff=0.8)
                if perto:
                    escolha, como = norm[perto[0]], "sugestao"
        parecidas = [norm[x] for x in difflib.get_close_matches(k, list(norm), n=3, cutoff=0.72)
                     if norm[x] != escolha]
        saida.append({"campo": campo, "coluna": escolha, "como": como, "sumiu": sumiu,
                      "parecidas": parecidas,
                      "novo": tem_salvo and k not in salvo_col and k not in nao_usar})
    return saida


def precisa_colunas(b: Base, an: Analise, sugestoes: list[dict]) -> str:
    """Por que a tela de colunas precisa abrir -- vazio quando não precisa.
    Abre na primeira carga de cada pasta e, depois, só quando algo mudou."""
    if not b.mapeavel:
        return ""
    salvo = mapeamento_salvo(b)
    if not salvo:
        return "primeira carga desta pasta"
    falta = [s["campo"].rotulo for s in sugestoes if s["campo"].obrigatorio and not s["coluna"]]
    if falta:
        return "campo obrigatório sem coluna: " + ", ".join(falta)
    if any(s["sumiu"] for s in sugestoes):
        return "uma coluna do mapeamento salvo não está no arquivo"
    if any(s["como"] == "sugestao" for s in sugestoes):
        return "coluna com nome parecido, para conferir"
    antes = salvo.get("colunas_do_arquivo")
    if antes is not None:
        agora = {normaliza(c) for c in an.colunas if c}
        antes_n = {normaliza(c) for c in antes}
        if agora - antes_n:
            n = len(agora - antes_n)
            return f"{n} coluna nova no arquivo" if n == 1 else f"{n} colunas novas no arquivo"
        if antes_n - agora:
            n = len(antes_n - agora)
            return f"{n} coluna saiu do arquivo" if n == 1 else f"{n} colunas saíram do arquivo"
    if salvo.get("aba_principal") and salvo["aba_principal"] != an.aba:
        return "a aba com os dados mudou"
    cab_salvo = ((salvo.get("abas") or {}).get(an.aba) or {}).get("cabecalho")
    if cab_salvo and int(cab_salvo) != an.cabecalho:
        return "a linha do cabeçalho mudou"
    return ""


def problemas(b: Base, escolhas: dict) -> list[str]:
    """O que impede de seguir: campo obrigatório sem coluna e coluna usada
    em dois campos (o mapeamento é de coluna para campo, um para um)."""
    saida = [f"{c.rotulo}: escolha a coluna (campo obrigatório)."
             for c in b.campos if c.obrigatorio and not escolhas.get(c.padrao)]
    usos: dict[str, list] = {}
    for c in b.campos:
        col = escolhas.get(c.padrao)
        if col:
            usos.setdefault(col, []).append(c.rotulo)
    saida += [f'A coluna "{col}" está em dois campos ({" e ".join(rot)}).'
              for col, rot in usos.items() if len(rot) > 1]
    return saida


def montar_mapeamento(b: Base, an: Analise, escolhas: dict, usuario: str = "") -> dict:
    """O JSON que o pipeline lê. `escolhas` é campo padrão -> coluna do
    arquivo (ou None para "não usar").

    "Não usar" num campo cujo nome padrão existe no arquivo precisa tirar essa
    coluna do caminho -- senão o pipeline a acharia pelo nome e usaria mesmo
    assim. Ela é renomeada para "(ignorada)" e o campo fica em nao_usar, para
    a próxima carga não sugerir de novo."""
    colunas: dict[str, str] = {}
    for padrao, col in escolhas.items():
        if col:
            colunas[col] = padrao
    usadas = set(colunas)
    norm: dict[str, str] = {}
    for c in an.colunas:
        if c:
            norm.setdefault(normaliza(c), c)
    nao_usar = []
    for campo in b.campos:
        if escolhas.get(campo.padrao):
            continue
        for nome in (campo.padrao, *campo.apelidos):
            c = norm.get(normaliza(nome))
            if c and c not in usadas:
                colunas[c] = f"{c} (ignorada)"
                usadas.add(c)
                if campo.padrao not in nao_usar:
                    nao_usar.append(campo.padrao)
    return {
        "base": b.codigo, "arquivo": b.arquivo,
        "atualizado_em": dt.datetime.now().isoformat(timespec="seconds"),
        "por": usuario,
        "aba_principal": an.aba,
        "abas": {an.aba: {"cabecalho": an.cabecalho, "colunas": colunas}},
        "nao_usar": nao_usar,
        "colunas_do_arquivo": [c for c in an.colunas if c],
    }


# ------------------------------------------------------------- comparação
def _chaves(b: Base, conteudo: bytes, an: Analise, escolhas: dict) -> set:
    import openpyxl
    cols = [escolhas.get(p) for p in b.chave]
    if not cols or not all(cols) or not all(c in an.colunas for c in cols):
        return set()
    idx = [an.colunas.index(c) for c in cols]
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    try:
        saida = set()
        for row in wb[an.aba].iter_rows(min_row=an.cabecalho + 1, values_only=True):
            chave = tuple(_texto(row[i]) if i < len(row) else "" for i in idx)
            if any(chave):
                saida.add(chave)
        return saida
    finally:
        wb.close()


def comparar(b: Base, conteudo_novo: bytes, an_novo: Analise, escolhas_novo: dict) -> dict:
    """A versão nova contra a que está em uso: linhas e, quando a base tem
    chave, quantos registros entram e quantos saem."""
    atual = arquivo_atual(b)
    res = {"linhas_atual": None, "linhas_novo": an_novo.linhas, "novos": None, "removidos": None}
    if not atual:
        return res
    conteudo_atual = atual.read_bytes()
    an_atual = analisar(b, conteudo_atual)
    res["linhas_atual"] = an_atual.linhas
    if b.chave and b.mapeavel and not an_atual.erro:
        esc_atual = {s["campo"].padrao: s["coluna"] for s in sugerir(b, an_atual.colunas)}
        k_novo = _chaves(b, conteudo_novo, an_novo, escolhas_novo)
        k_atual = _chaves(b, conteudo_atual, an_atual, esc_atual)
        if k_novo and k_atual:
            res["novos"] = len(k_novo - k_atual)
            res["removidos"] = len(k_atual - k_novo)
    return res


# ------------------------------------------------------------- execução
class Ocupado(RuntimeError):
    """Já tem uma atualização rodando neste Gplan."""


class FalhaNaAtualizacao(RuntimeError):
    """Um dos scripts parou com erro."""


@dataclass
class Tarefa:
    titulo: str
    etapas: list                    # [(chave, o que faz, detalhe, tempo esperado)]
    id: str = field(default_factory=lambda: uuid.uuid4().hex[:10])
    inicio: dt.datetime = field(default_factory=dt.datetime.now)
    etapa: str = ""
    tempos: dict = field(default_factory=dict)      # etapa concluída -> segundos
    falhas: dict = field(default_factory=dict)      # etapa -> o que deu errado
    linhas: list = field(default_factory=list)
    fim: dt.datetime | None = None
    erro: str = ""
    avisos: list = field(default_factory=list)
    resultado: dict = field(default_factory=dict)
    _t0: float = 0.0

    @property
    def rodando(self) -> bool:
        return self.fim is None

    def passo(self, chave: str) -> None:
        agora = time.monotonic()
        if self.etapa and self.etapa not in self.falhas:
            self.tempos[self.etapa] = agora - self._t0
        self.etapa, self._t0 = chave, agora

    def falhou(self, chave: str, motivo: str) -> None:
        self.falhas[chave] = motivo

    def escreve(self, linha: str) -> None:
        linha = linha.rstrip()
        if linha:
            self.linhas.append(linha)
            if len(self.linhas) > 600:
                del self.linhas[:-400]

    def em_andamento(self) -> float:
        return time.monotonic() - self._t0 if self._t0 else 0.0

    def segundos(self) -> float:
        return ((self.fim or dt.datetime.now()) - self.inicio).total_seconds()

    def terminou_ha(self) -> float:
        """Segundos desde o fim (0 enquanto roda)."""
        return (dt.datetime.now() - self.fim).total_seconds() if self.fim else 0.0


_TRAVA = threading.Lock()
_ATUAL: Tarefa | None = None


def tarefa_atual() -> Tarefa | None:
    return _ATUAL


_QUEM = {"nome": ""}


def quem_esta_usando(nome: str) -> None:
    """Quem abriu a aba Bases -- a tela avisa a cada desenho. É o nome que
    entra no registro das atualizações."""
    _QUEM["nome"] = nome or ""


def arquivo_registro() -> Path:
    return historico_dir() / "registro.jsonl"


def anotar(tarefa: Tarefa) -> None:
    """Uma linha por atualização terminada, no fim do registro.

    Uma linha por vez, em JSON: cresce ~200 bytes por atualização, abre em
    qualquer editor e não se perde se o Gplan for fechado no meio.
    """
    linha = {
        "quando": tarefa.inicio.isoformat(timespec="seconds"),
        "quem": _QUEM["nome"],
        "o_que": tarefa.titulo,
        "base": tarefa.resultado.get("base", ""),
        "modo": tarefa.resultado.get("modo", ""),
        "arquivo": tarefa.resultado.get("arquivo", ""),
        "segundos": round(tarefa.segundos()),
        "resultado": ("erro" if tarefa.erro else "aviso" if tarefa.avisos else "ok"),
        "publicacao": tarefa.resultado.get("publicacao", ""),
        "detalhe": tarefa.erro or " | ".join(tarefa.avisos),
    }
    try:
        arq = arquivo_registro()
        arq.parent.mkdir(parents=True, exist_ok=True)
        with open(arq, "a", encoding="utf-8") as f:
            f.write(json.dumps(linha, ensure_ascii=False) + chr(10))
    except OSError:
        pass        # registro é apoio: falta dele não derruba a atualização


def registro(limite: int = 50) -> list:
    """As últimas atualizações, da mais nova para a mais velha."""
    arq = arquivo_registro()
    if not arq.exists():
        return []
    linhas = []
    for bruta in arq.read_text(encoding="utf-8").splitlines()[-limite:]:
        try:
            linhas.append(json.loads(bruta))
        except ValueError:
            continue
    return list(reversed(linhas))


def _disparar(tarefa: Tarefa, trabalho) -> Tarefa:
    global _ATUAL
    if not _TRAVA.acquire(blocking=False):
        raise Ocupado("Já tem uma atualização rodando. Espere ela terminar.")
    _ATUAL = tarefa

    def corre():
        try:
            trabalho(tarefa)
        except Exception as erro:
            tarefa.erro = str(erro) or type(erro).__name__
            if tarefa.etapa:
                tarefa.falhou(tarefa.etapa, tarefa.erro)
        finally:
            if tarefa.etapa and tarefa.etapa not in tarefa.falhas:
                tarefa.tempos[tarefa.etapa] = tarefa.em_andamento()
            tarefa.fim = dt.datetime.now()
            anotar(tarefa)
            _TRAVA.release()

    threading.Thread(target=corre, name=f"gplan-bases-{tarefa.id}", daemon=True).start()
    return tarefa


_PREFIXO_LOG = re.compile(r"^\[[^\]]+\]\s+(\w+)\s+-\s+[\w.]+\s+-\s?")

# O RECALCULAR_BASES diz numa linha o que houve com cada base; o fim da saída
# só repete "feche no Excel", o que engana quando o Excel nem abriu o arquivo.
_PULADA = re.compile(r"^\s*PULADA\s+(.+?)\s+--")
_FALHOU = re.compile(r"^\s*abrindo\s+(.+?)\s+->\s+FALHOU:\s*(.*)$")


def _problema_da_base(linha: str) -> str:
    m = _PULADA.match(linha)
    if m:
        return (f"{m.group(1)} está aberta no Excel (ou o OneDrive está sincronizando): "
                "feche e tente de novo")
    m = _FALHOU.match(linha)
    if not m:
        return ""
    arquivo, motivo = m.groups()
    if "open" in motivo.lower() and "workbooks" in motivo.lower():
        return (f"o Excel não conseguiu abrir {arquivo}: o arquivo pode estar corrompido ou "
                "num formato que o Excel não lê. Abra o arquivo no Excel, salve de novo e "
                "tente outra vez")
    return f"o Excel falhou em {arquivo}: {motivo}"


def _rodar(t: Tarefa, cmd: list[str], cwd: Path, so_log: bool = True,
           encoding: str = "utf-8") -> None:
    """Roda um script e passa para a tarefa o que ele registra. Dos scripts
    em Python fica só o que é linha de log (o resumo em JSON do fim é ruído
    aqui); do PowerShell, tudo."""
    env = os.environ.copy()
    env.setdefault("GPLAN_BASES_DIR", str(bases_dir()))
    env["PYTHONIOENCODING"] = "utf-8"
    # sem isto o Python do script guarda a saída em bloco (é um pipe, não
    # um terminal) e a tela só via o registro no fim
    env["PYTHONUNBUFFERED"] = "1"
    proc = subprocess.Popen(cmd, cwd=str(cwd), stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                            text=True, encoding=encoding, errors="replace", env=env,
                            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    ultimo_erro = ""
    fim = []                    # as últimas linhas, para explicar uma falha
    problemas = []              # base a base, do RECALCULAR_BASES
    for bruta in proc.stdout:
        if bruta.strip():
            fim = (fim + [bruta.strip()])[-4:]
            problema = _problema_da_base(bruta.strip())
            if problema:
                problemas.append(problema)
        m = _PREFIXO_LOG.match(bruta)
        if m:
            texto = bruta[m.end():]
            if m.group(1) in ("ERROR", "CRITICAL"):
                ultimo_erro = texto.strip()
            t.escreve(texto)
        elif not so_log:
            t.escreve(bruta)
    codigo = proc.wait()
    if codigo != 0:
        raise FalhaNaAtualizacao("; ".join(problemas)
                                 or _explicar(ultimo_erro or " / ".join(fim))
                                 or f"o script parou com o código {codigo}")


def _explicar(erro: str) -> str:
    if "Permission denied" in erro or "PermissionError" in erro or "WinError 32" in erro:
        return (erro + " -- a planilha ou a base está aberta no Excel (ou o OneDrive está "
                "sincronizando). Feche e tente de novo.")
    return erro


def _importar(t: Tarefa, planilha: Path, codigos: list[str] | None) -> None:
    cmd = [sys.executable, "update_current_workbook_from_bases.py", str(planilha)]
    if codigos:
        cmd += ["--bases", ",".join(codigos)]
    _rodar(t, cmd, ferramentas_dir())


def _resumo(t: Tarefa, planilha: Path) -> None:
    _rodar(t, [sys.executable, "apply_excel_safe_control_formulas.py", str(planilha)],
           ferramentas_dir())


def _recalcular(t: Tarefa, codigos: list[str] | None) -> None:
    cmd = ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass",
           "-File", str(recalcular_ps1())]
    if codigos:
        cmd += ["-Somente", ",".join(codigos)]
    # o PowerShell escreve no código de página do console (ibm850 num Windows
    # em português), e não em UTF-8
    _rodar(t, cmd, controle_dir(), so_log=False, encoding="oem")


def etapas_da_carga(b: Base, recalcular: bool, publica: bool, trocar: bool = True) -> list:
    """(chave, o que faz, detalhe, tempo esperado): a mesma lista que a tela
    mostra antes de aplicar e que acompanha a execução. Os tempos são os
    medidos em 11/09/2026: importar uma base, ~1 min; o resumo, de 45 s a
    3 min (mais lento na pasta do OneDrive)."""
    e = []
    if trocar:
        e.append(("guardar", "Guarda a versão em uso no histórico",
                  f"03_HISTORICO_IMPORTACOES › {b.stem}", "na hora"))
        e.append(("trocar", "Troca o arquivo na pasta das bases",
                  f"00_COLOCAR_ATUALIZADAS_AQUI › {b.stem}", "na hora"))
    if recalcular:
        e.append(("excel", "Recalcula no Excel",
                  "abre, recalcula e salva" + (" · atualiza o vínculo do SharePoint"
                                               if b.codigo == "01" else ""), "segundos"))
    e.append(("importar", f"Importa a base {b.codigo}", " · ".join(b.refeitas), "~1 min"))
    if b.precisa_resumo:
        e.append(("resumo", "Refaz o resumo por TAG", RESUMO_ABAS, "1 a 3 min"))
    if publica:
        e.append(("confere", "Confere o que mudou na planilha",
                  "linhas por aba, antes e depois", "~20 s"))
    e.append(_etapa_publicar(publica))
    return e


def _etapa_publicar(publica: bool) -> tuple:
    if publica:
        return ("publicar", "Publica no Supabase", PLANILHA_SUPABASE, "automático")
    return ("publicar", "Publicação no Supabase desligada",
            "esta instância lê a planilha do disco", "—")


def _nome_versao(arquivo: Path, pasta: Path) -> Path:
    quando = dt.datetime.fromtimestamp(arquivo.stat().st_mtime).strftime("%Y-%m-%d_%H%M")
    alvo = pasta / f"{quando}{arquivo.suffix}"
    n = 2
    while alvo.exists():
        alvo = pasta / f"{quando}_{n}{arquivo.suffix}"
        n += 1
    return alvo


def _podar(b: Base) -> list[Path]:
    """Mantém as VERSOES_GUARDADAS mais recentes da pasta da base."""
    apagadas = []
    for velho in versoes(b)[VERSOES_GUARDADAS:]:
        velho.unlink()
        apagadas.append(velho)
    return apagadas


def _trocar(t: Tarefa, b: Base, conteudo: bytes, ext: str, mapeamento: dict | None) -> dict:
    """Passos 1 e 2: a versão em uso vai para o histórico e a nova entra no
    lugar, com o mapeamento dela. Devolve o que é preciso para desfazer."""
    entrada = entrada_dir()
    pasta_hist = historico_dir() / b.stem
    pasta_hist.mkdir(parents=True, exist_ok=True)
    arq_map = mapeamentos_dir() / f"{b.stem}.json"
    troca = {"guardadas": [], "destino": None, "arq_map": arq_map,
             "map_antes": arq_map.read_bytes() if arq_map.exists() else None}
    try:
        # todas as variantes em uso saem: a 07 pode ter .xlsm e .xlsx, e se
        # ficasse o outro o pipeline leria o arquivo errado
        por_nome = {_nome(p.name): p for p in entrada.iterdir() if p.is_file()}
        for e in b.extensoes:
            atual = por_nome.get(_nome(b.stem + e))
            if atual:
                destino_hist = _nome_versao(atual, pasta_hist)
                shutil.move(str(atual), str(destino_hist))
                troca["guardadas"].append((atual, destino_hist))
        t.passo("trocar")
        destino = entrada / f"{b.stem}{ext}"
        for original, _ in troca["guardadas"]:
            if original.suffix.lower() == ext:
                destino = original          # mantém o nome exato de hoje (a 04 é .XLSX)
        tmp = destino.with_name(destino.name + ".tmp")
        tmp.write_bytes(conteudo)
        os.replace(tmp, destino)
        troca["destino"] = destino
        if mapeamento is not None:
            arq_map.parent.mkdir(parents=True, exist_ok=True)
            arq_map.write_text(json.dumps(mapeamento, ensure_ascii=False, indent=2),
                               encoding="utf-8")
    except PermissionError as erro:
        _desfazer(troca)
        raise FalhaNaAtualizacao(
            f"{Path(erro.filename).name if erro.filename else b.arquivo} está aberto no Excel "
            "(ou o OneDrive está sincronizando). Feche e tente de novo.") from erro
    except Exception:
        _desfazer(troca)
        raise
    return troca


def _desfazer(troca: dict) -> None:
    if troca["destino"] is not None:
        troca["destino"].unlink(missing_ok=True)
    for original, hist in reversed(troca["guardadas"]):
        shutil.move(str(hist), str(original))
    arq = troca["arq_map"]
    if troca["map_antes"] is None:
        arq.unlink(missing_ok=True)
    else:
        arq.write_bytes(troca["map_antes"])


# Quanto uma aba pode encolher sem ser suspeito. Abaixo disso é variação
# normal da base; acima, quase sempre é arquivo errado ou coluna fora do lugar.
ENCOLHEU_FRACAO = 0.10
ENCOLHEU_MINIMO = 20


def _encolheu(antes: dict, depois: dict) -> list:
    """As abas que perderam linhas demais entre antes e depois da importação."""
    perdas = []
    for aba, tinha in (antes or {}).items():
        agora = (depois or {}).get(aba)
        if agora is None or not tinha:
            continue
        perda = tinha - agora
        if perda >= ENCOLHEU_MINIMO and perda >= tinha * ENCOLHEU_FRACAO:
            perdas.append(f"{aba} caiu de {tinha} para {agora} linhas")
    return perdas


def _depois_de_importar(t: Tarefa, resumo: bool, planilha: Path, publicar,
                        antes: dict | None = None) -> None:
    """Resumo por TAG e publicação. A base já entrou: o que falhar daqui para
    frente vira aviso com o que fazer, e não desfaz a carga."""
    if resumo:
        t.passo("resumo")
        try:
            _resumo(t, planilha)
        except Exception as erro:
            t.falhou("resumo", str(erro))
            t.resultado["falta_resumo"] = True
            t.avisos.append("A base entrou, mas o resumo por TAG não foi refeito, e por isso a "
                            f"planilha não foi publicada. Motivo: {erro}")
            return
    if antes and publicar is not None:
        t.passo("confere")
        perdas = _encolheu(antes, contagens(planilha))
        if perdas:
            t.falhou("confere", "; ".join(perdas))
            t.resultado["falta_publicar"] = True
            t.avisos.append(
                "A planilha foi atualizada aqui, mas NÃO subiu para o Supabase: "
                + "; ".join(perdas) + ". Confira se isso era esperado -- se for, "
                "clique em Publicar de novo.")
            return
    t.passo("publicar")
    if publicar is None:
        t.resultado["publicacao"] = "desligada"
        return
    try:
        t.resultado["publicacao"] = publicar(planilha)
    except Exception as erro:
        t.falhou("publicar", str(erro))
        t.resultado["falta_publicar"] = True
        t.avisos.append("A planilha foi atualizada aqui, mas não subiu para o Supabase "
                        f"({type(erro).__name__}: {erro}).")


def iniciar_carga(b: Base, conteudo: bytes, extensao: str, mapeamento: dict | None,
                  planilha: Path, recalcular: bool, publicar=None) -> Tarefa:
    """Troca a base na pasta e refaz o que depende dela, numa thread.

    Se o Excel ou a importação falharem, a pasta volta ao que era -- a base
    anterior e o mapeamento anterior -- e a planilha fica como estava (o
    pipeline só grava no fim, com tudo lido)."""
    ext = extensao.lower()
    if ext not in b.extensoes:
        raise ValueError(f"A pasta {b.codigo} espera {' ou '.join(b.extensoes)}; "
                         f"o arquivo é {ext}.")
    t = Tarefa(titulo=f"{b.codigo} · {b.nome}",
               etapas=etapas_da_carga(b, recalcular, publicar is not None))
    t.resultado.update(base=b.codigo, modo="carga")
    # as linhas de cada aba antes de mexer: é com isto que a conferência
    # de depois compara (já está em cache, a lista das bases acabou de ler)
    antes = contagens(planilha)

    def trabalho(t: Tarefa):
        t.passo("guardar")
        troca = _trocar(t, b, conteudo, ext, mapeamento)
        try:
            if recalcular:
                t.passo("excel")
                _recalcular(t, [b.codigo])
            t.passo("importar")
            _importar(t, planilha, [b.codigo])
        except BaseException:
            _desfazer(troca)
            t.resultado["desfeito"] = True
            raise
        t.resultado["guardadas"] = [h.name for _, h in troca["guardadas"]]
        t.resultado["apagadas"] = [p.name for p in _podar(b)]
        _depois_de_importar(t, b.precisa_resumo, planilha, publicar, antes)

    return _disparar(t, trabalho)


def iniciar_reprocessar(b: Base, planilha: Path, recalcular: bool, publicar=None) -> Tarefa:
    """Importa a versão que já está na pasta -- para quando o arquivo foi
    editado direto nela (o Rundown é preenchido assim)."""
    t = Tarefa(titulo=f"{b.codigo} · {b.nome}",
               etapas=etapas_da_carga(b, recalcular, publicar is not None, trocar=False))
    t.resultado.update(base=b.codigo, modo="reprocessar")
    # as linhas de cada aba antes de mexer: é com isto que a conferência
    # de depois compara (já está em cache, a lista das bases acabou de ler)
    antes = contagens(planilha)

    def trabalho(t: Tarefa):
        if recalcular:
            t.passo("excel")
            _recalcular(t, [b.codigo])
        t.passo("importar")
        _importar(t, planilha, [b.codigo])
        _depois_de_importar(t, b.precisa_resumo, planilha, publicar, antes)

    return _disparar(t, trabalho)


def etapas_de_tudo(publica: bool) -> list:
    e = [("excel", "Recalcula as 11 bases no Excel", "abre, recalcula e salva cada uma",
             "alguns min"),
            ("importar", "Importa as 11 bases", "todas as abas de base do controle",
             "2 a 5 min"),
            ("resumo", "Refaz o resumo por TAG", RESUMO_ABAS, "1 a 3 min")]
    if publica:
        e.append(("confere", "Confere o que mudou na planilha",
                  "linhas por aba, antes e depois", "~20 s"))
    e.append(_etapa_publicar(publica))
    return e


def iniciar_tudo(planilha: Path, publicar=None) -> Tarefa:
    """O ATUALIZAR_TUDO.cmd, os três passos na mesma ordem, mais a publicação."""
    t = Tarefa(titulo="Atualizar tudo", etapas=etapas_de_tudo(publicar is not None))
    t.resultado.update(modo="tudo")
    # as linhas de cada aba antes de mexer: é com isto que a conferência
    # de depois compara (já está em cache, a lista das bases acabou de ler)
    antes = contagens(planilha)

    def trabalho(t: Tarefa):
        t.passo("excel")
        _recalcular(t, None)
        t.passo("importar")
        _importar(t, planilha, None)
        _depois_de_importar(t, True, planilha, publicar, antes)

    return _disparar(t, trabalho)


def iniciar_resumo(planilha: Path, publicar=None) -> Tarefa:
    """Refaz só o passo 3 e publica: o que falta quando o resumo falhou."""
    t = Tarefa(titulo="Resumo por TAG",
               etapas=[("resumo", "Refaz o resumo por TAG", RESUMO_ABAS, "1 a 3 min"),
                       _etapa_publicar(publicar is not None)])
    t.resultado.update(modo="resumo")
    return _disparar(t, lambda t: _depois_de_importar(t, True, planilha, publicar))


def iniciar_publicacao(planilha: Path, publicar) -> Tarefa:
    t = Tarefa(titulo="Publicação", etapas=[_etapa_publicar(True)])
    t.resultado.update(modo="publicar")
    return _disparar(t, lambda t: _depois_de_importar(t, False, planilha, publicar))
